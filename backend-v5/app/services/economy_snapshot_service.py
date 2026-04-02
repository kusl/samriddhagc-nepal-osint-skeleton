"""NRB workbook-backed economy snapshot service."""

from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Optional

from app.services.debt_clock_service import DebtClockService


try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - handled at runtime in prod image
    load_workbook = None


MetricResolver = Callable[["EconomySnapshotService"], tuple[Optional[float], Optional[str]]]


@dataclass(frozen=True)
class MetricSpec:
    key: str
    label: str
    unit: str
    meta: str
    resolver: MetricResolver


@dataclass(frozen=True)
class SectionSpec:
    key: str
    label: str
    badge: str
    metrics: tuple[MetricSpec, ...]


class EconomySnapshotService:
    """Build a reusable economy snapshot from official NRB workbook data."""

    CACHE_KEY = "economy:nrb_snapshot:v1"
    CACHE_TTL_SECONDS = 6 * 60 * 60
    WORKBOOK_PATH = Path(__file__).resolve().parents[1] / "data" / "nrb" / "current_macro_2082_83_seven_months.xlsx"
    WORKBOOK_PERIOD_LABEL = "Seven months of FY 2082/83"
    WORKBOOK_AS_OF_LABEL = "Magh 2082 (Mid-February, 2026)"

    def __init__(self, redis: Any):
        self.redis = redis
        self._workbook = None
        self._sheet_cache: dict[str, list[tuple[Any, ...]]] = {}
        self._debt_summary: dict[str, Any] | None = None

    async def get_snapshot(self, force_refresh: bool = False) -> dict[str, Any]:
        workbook_mtime = self._workbook_mtime()
        if not force_refresh:
            cached = await self.redis.get(self.CACHE_KEY)
            if cached:
                payload = json.loads(cached)
                if payload.get("workbook_mtime") == workbook_mtime:
                    return payload

        payload = await self._build_snapshot()
        payload["workbook_mtime"] = workbook_mtime
        await self.redis.set(self.CACHE_KEY, json.dumps(payload), ex=self.CACHE_TTL_SECONDS)
        return payload

    async def _build_snapshot(self) -> dict[str, Any]:
        self._debt_summary = await DebtClockService(self.redis).get_nepal_summary()
        extracted_at = datetime.now(timezone.utc).isoformat()
        sections = {
            spec.key: {
                "key": spec.key,
                "label": spec.label,
                "badge": spec.badge,
                "as_of_label": self.WORKBOOK_AS_OF_LABEL,
                "metrics": [self._render_metric(metric) for metric in spec.metrics],
            }
            for spec in self._sections()
        }
        return {
            "source_label": "Nepal Rastra Bank workbook + official macro cache",
            "workbook_label": self.WORKBOOK_PATH.name,
            "workbook_period_label": self.WORKBOOK_PERIOD_LABEL,
            "as_of_label": self.WORKBOOK_AS_OF_LABEL,
            "extracted_at": extracted_at,
            "sections": sections,
        }

    def _render_metric(self, spec: MetricSpec) -> dict[str, Any]:
        raw_value, source_label = spec.resolver(self)
        return {
            "key": spec.key,
            "label": spec.label,
            "unit": spec.unit,
            "raw_value": raw_value,
            "display_value": self._format_value(raw_value, spec.unit),
            "meta": spec.meta,
            "source_label": source_label or self.WORKBOOK_AS_OF_LABEL,
        }

    def _sections(self) -> tuple[SectionSpec, ...]:
        return (
            SectionSpec(
                key="fiscal_position",
                label="Fiscal Position",
                badge="NRB",
                metrics=(
                    MetricSpec("total_expenditure", "Total Expenditure", "npr_million", "Government budgetary operations, seven months", self._from_row("34.GBO", "Total Expenditure")),
                    MetricSpec("total_revenue", "Total Revenue", "npr_million", "Government revenue collection, seven months", self._from_row_nth_numeric("35.Revenue", "Total  Revenue", 5)),
                    MetricSpec("total_resources", "Total Resources", "npr_million", "Resources available in NRB budget operations table", self._from_row("34.GBO", "Total Resources [2.1+2.2]")),
                    MetricSpec("resource_gap", "Resource Gap", "npr_million", "Expenditure minus total resources, seven months", self._resource_gap),
                    MetricSpec("customs_revenue", "Customs Revenue", "npr_million", "Customs revenue collected in seven months", self._from_row_nth_numeric("35.Revenue", "Customs", 5)),
                    MetricSpec("revenue_coverage", "Revenue Coverage", "pct", "Total revenue as a share of total expenditure", self._revenue_coverage),
                ),
            ),
            SectionSpec(
                key="external_sector",
                label="External Sector",
                badge="NRB",
                metrics=(
                    MetricSpec("exports", "Exports", "npr_million", "Total exports, seven months", self._from_row_nth_numeric("13.Direction", "TOTAL EXPORTS", 3)),
                    MetricSpec("imports", "Imports", "npr_million", "Total imports, seven months", self._from_row_nth_numeric("13.Direction", "TOTAL IMPORTS", 3)),
                    MetricSpec("trade_balance", "Trade Balance", "npr_million_signed", "Trade balance, seven months", self._from_row_nth_numeric("13.Direction", "TOTAL TRADE BALANCE", 3)),
                    MetricSpec("current_account", "Current Account", "npr_million_signed", "Balance of payments current account, seven months", self._from_bop_row("Current account")),
                    MetricSpec("remittances", "Remittances", "npr_million", "Workers' remittances, seven months", self._from_bop_row("Workers' remittances")),
                    MetricSpec("fx_reserves_usd", "FX Reserves", "usd_million", "Gross foreign exchange reserves, mid-February", self._from_row_nth_numeric("31.Reserves $", "C. Gross Foreign Exchange Reserves", 3)),
                ),
            ),
            SectionSpec(
                key="monetary_conditions",
                label="Monetary Conditions",
                badge="NRB",
                metrics=(
                    MetricSpec("broad_money_growth", "Broad Money Growth", "pct", "Latest official broad money growth", self._from_debt_field("broad_money_growth_pct")),
                    MetricSpec("private_credit_growth", "Private Credit Growth", "pct", "Latest official private-sector credit growth", self._from_debt_field("private_sector_credit_growth_pct")),
                    MetricSpec("net_foreign_assets", "Net Foreign Assets", "npr_million", "Monetary survey, mid-February", self._from_row_cell("37.MS", "1. Foreign Assets, Net", 4)),
                    MetricSpec("claims_on_government", "Claims on Government", "npr_million", "Net claims on government, mid-February", self._from_row_cell("37.MS", "a. Net Claims on Government", 4)),
                    MetricSpec("net_liquidity", "Net Liquidity", "npr_million_signed", "Net liquidity injection (+) / absorption (-)", self._from_row("52.MO Summary", "C. Net Liquidity Injection (+) / Absorption (-)")),
                    MetricSpec("banking_network", "Banking Network", "network", "Licensed BFIs and branch footprint", self._banking_network),
                ),
            ),
            SectionSpec(
                key="prices_cost_pressure",
                label="Prices & Cost Pressure",
                badge="NRB",
                metrics=(
                    MetricSpec("cpi_inflation", "Inflation", "pct", "Headline CPI, year-on-year", self._from_debt_field("inflation_pct")),
                    MetricSpec("food_inflation", "Food Inflation", "pct", "Food CPI, year-on-year", self._from_debt_field("food_inflation_pct")),
                    MetricSpec("non_food_inflation", "Non-food Inflation", "pct", "Non-food CPI, year-on-year", self._from_debt_field("non_food_inflation_pct")),
                    MetricSpec("deposit_rate", "Deposit Rate", "pct", "Weighted average commercial-bank deposit rate", self._from_interest_row("Weighted Average Deposit Rate")),
                    MetricSpec("lending_rate", "Lending Rate", "pct", "Weighted average commercial-bank lending rate", self._from_interest_row("Weighted Average Lending Rate")),
                    MetricSpec("bank_rate", "Bank Rate", "pct", "NRB bank rate", self._from_interest_row("Bank Rate")),
                ),
            ),
        )

    def _workbook_mtime(self) -> float:
        return self.WORKBOOK_PATH.stat().st_mtime

    def _get_workbook(self):
        if load_workbook is None:
            raise RuntimeError("openpyxl is not installed")
        if self._workbook is None:
            self._workbook = load_workbook(self.WORKBOOK_PATH, data_only=True, read_only=True)
        return self._workbook

    def _rows(self, sheet_name: str) -> list[tuple[Any, ...]]:
        if sheet_name not in self._sheet_cache:
            sheet = self._get_workbook()[sheet_name]
            self._sheet_cache[sheet_name] = [tuple(row) for row in sheet.iter_rows(values_only=True)]
        return self._sheet_cache[sheet_name]

    @staticmethod
    def _normalize(value: Any) -> str:
        if value is None:
            return ""
        return " ".join(str(value).strip().lower().replace("\n", " ").split())

    def _find_row(self, sheet_name: str, needle: str) -> tuple[Any, ...]:
        normalized_needle = self._normalize(needle)
        for row in self._rows(sheet_name):
            haystack = " ".join(self._normalize(cell) for cell in row[:4] if cell is not None)
            if normalized_needle in haystack:
                return row
        raise KeyError(f"Row not found for {needle!r} in {sheet_name}")

    @staticmethod
    def _latest_numeric(row: tuple[Any, ...]) -> Optional[float]:
        for cell in reversed(row):
            if isinstance(cell, (int, float)):
                return float(cell)
        return None

    def _from_row(self, sheet_name: str, needle: str) -> MetricResolver:
        def resolver(service: "EconomySnapshotService") -> tuple[Optional[float], Optional[str]]:
            row = service._find_row(sheet_name, needle)
            return service._latest_numeric(row), service.WORKBOOK_AS_OF_LABEL

        return resolver

    def _from_bop_row(self, needle: str) -> MetricResolver:
        def resolver(service: "EconomySnapshotService") -> tuple[Optional[float], Optional[str]]:
            row = service._find_row("28(A).BoP_Cumulative", needle)
            numbers = [float(value) for value in row if isinstance(value, (int, float))]
            return (numbers[-1] if numbers else None), service.WORKBOOK_AS_OF_LABEL

        return resolver

    def _from_row_nth_numeric(self, sheet_name: str, needle: str, nth_from_end: int) -> MetricResolver:
        def resolver(service: "EconomySnapshotService") -> tuple[Optional[float], Optional[str]]:
            row = service._find_row(sheet_name, needle)
            numbers = [float(value) for value in row if isinstance(value, (int, float))]
            if len(numbers) < nth_from_end:
                return None, service.WORKBOOK_AS_OF_LABEL
            return numbers[-nth_from_end], service.WORKBOOK_AS_OF_LABEL

        return resolver

    def _from_row_cell(self, sheet_name: str, needle: str, cell_index: int) -> MetricResolver:
        def resolver(service: "EconomySnapshotService") -> tuple[Optional[float], Optional[str]]:
            row = service._find_row(sheet_name, needle)
            value = row[cell_index] if len(row) > cell_index else None
            return (float(value) if isinstance(value, (int, float)) else None), service.WORKBOOK_AS_OF_LABEL

        return resolver

    def _from_interest_row(self, needle: str) -> MetricResolver:
        def resolver(service: "EconomySnapshotService") -> tuple[Optional[float], Optional[str]]:
            row = service._find_row("55.Interest Rate", needle)
            return service._latest_numeric(row), service.WORKBOOK_AS_OF_LABEL

        return resolver

    def _from_debt_field(self, field: str) -> MetricResolver:
        def resolver(service: "EconomySnapshotService") -> tuple[Optional[float], Optional[str]]:
            summary = service._debt_summary
            if summary is None:
                raise RuntimeError("Debt summary not loaded")
            label_key_overrides = {
                "inflation_pct": "inflation_label",
                "food_inflation_pct": "food_inflation_label",
                "non_food_inflation_pct": "non_food_inflation_label",
                "broad_money_growth_pct": "broad_money_growth_label",
                "private_sector_credit_growth_pct": "private_sector_credit_growth_label",
                "weighted_deposit_rate_pct": "weighted_deposit_rate_label",
                "weighted_credit_rate_pct": "weighted_credit_rate_label",
                "interbank_rate_pct": "interbank_rate_label",
            }
            label = summary.get(label_key_overrides.get(field, f"{field}_label"))
            value = summary.get(field)
            return (float(value) if value is not None else None), label or service.WORKBOOK_AS_OF_LABEL

        return resolver

    def _resource_gap(self, service: "EconomySnapshotService") -> tuple[Optional[float], Optional[str]]:
        expenditure, _ = self._from_row("34.GBO", "Total Expenditure")(service)
        resources, _ = self._from_row("34.GBO", "Total Resources [2.1+2.2]")(service)
        if expenditure is None or resources is None:
            return None, service.WORKBOOK_AS_OF_LABEL
        return expenditure - resources, service.WORKBOOK_AS_OF_LABEL

    def _revenue_coverage(self, service: "EconomySnapshotService") -> tuple[Optional[float], Optional[str]]:
        revenue, _ = self._from_row_nth_numeric("35.Revenue", "Total  Revenue", 5)(service)
        expenditure, _ = self._from_row("34.GBO", "Total Expenditure")(service)
        if revenue is None or expenditure in (None, 0):
            return None, service.WORKBOOK_AS_OF_LABEL
        return (revenue / expenditure) * 100, service.WORKBOOK_AS_OF_LABEL

    def _banking_network(self, service: "EconomySnapshotService") -> tuple[Optional[float], Optional[str]]:
        summary = service._debt_summary
        if summary is None:
            raise RuntimeError("Debt summary not loaded")
        institutions = summary.get("total_financial_institutions")
        branches = summary.get("total_bfi_branches")
        if institutions is None and branches is None:
            return None, service.WORKBOOK_AS_OF_LABEL
        packed = float(f"{int(institutions or 0)}.{int(branches or 0):05d}")
        return packed, summary.get("total_bfi_branches_label") or service.WORKBOOK_AS_OF_LABEL

    def _format_value(self, value: Optional[float], unit: str) -> str:
        if value is None:
            return "N/A"
        if unit == "pct":
            return f"{value:.2f}%"
        if unit == "npr_million":
            return self._format_compact(value, "Nrs", 1_000_000)
        if unit == "npr_million_signed":
            prefix = "-" if value < 0 else ""
            return f"{prefix}{self._format_compact(abs(value), 'Nrs', 1_000_000)}"
        if unit == "usd_million":
            return self._format_compact(value, "US$", 1_000_000)
        if unit == "network":
            integer = int(value)
            fractional = f"{value:.5f}".split(".")[1]
            branches = int(fractional)
            return f"{integer} FI / {branches:,} branches"
        return f"{value:,.2f}"

    @staticmethod
    def _format_compact(value_million: float, prefix: str, scale: float) -> str:
        absolute_value = value_million * scale
        if absolute_value >= 1_000_000_000_000:
            return f"{prefix} {(absolute_value / 1_000_000_000_000):.2f}T".replace(".00", "")
        if absolute_value >= 1_000_000_000:
            return f"{prefix} {(absolute_value / 1_000_000_000):.2f}B".replace(".00", "")
        if absolute_value >= 1_000_000:
            return f"{prefix} {(absolute_value / 1_000_000):.2f}M".replace(".00", "")
        return f"{prefix} {absolute_value:,.0f}"
